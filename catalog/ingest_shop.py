"""Import a permitted shop's CSV or first XLSX worksheet into item.

Run from catalog: uv run --with openpyxl ingest_shop.py --shop my-shop \
    --file products.xlsx --permission 'written permission 2026-09-26' --dry-run
XLSX requires openpyxl. CSV supports UTF-8 (including BOM), comma, semicolon
and tab delimiters. Dimensions are centimetres; prices round half up to whole
AMD. Permission text is retained exactly as supplied. No URLs are fetched.
"""
import argparse
from collections import Counter
import csv
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import os
from pathlib import Path
import re

import psycopg
from psycopg.types.json import Jsonb

from ingest_abo import kind_of


ALIASES = {
    "sku": "sku|article|article number|артикул|արտիկուլ|կոդ",
    "name": "name|product name|название|наименование|անվանում|անուն",
    "category": "category|product type|категория|тип|կատեգորիա|տեսակ",
    "width": "width|ширина|լայնություն",
    "depth": "depth|глубина|խորություն",
    "height": "height|высота|բարձրություն",
    "price": "price|цена|стоимость|գին",
    "color": "color|colour|цвет|գույն",
    "material": "material|materials|материал|материалы|նյութ",
    "photo": "photo|photos|photo url|photo urls|image|images|image url|image urls|"
             "фото|фотография|фотографии|ссылка на фото|изображение|изображения|նկար|նկարներ|լուսանկար",
    "url": "url|product url|product link|ссылка|ссылка на товар|товар url|հղում|ապրանքի հղում",
}
HEADER_FIELDS = {alias: field for field, aliases in ALIASES.items() for alias in aliases.split("|")}
COLUMNS = (
    "id", "source", "source_id", "name", "brand", "description", "product_type", "kind",
    "size_m", "fit_size_m", "size_status", "size_evidence", "listing_size_m", "price",
    "currency", "price_source", "color_text", "color_std", "materials", "styles", "keywords",
    "main_image_url", "image_urls", "glb_url", "license", "raw",
)
# Same ABO columns, with an explicit false flag on both insert and update.
UPSERT = (
    f"insert into item ({', '.join(COLUMNS)}, editor_set) "
    f"values ({', '.join(['%s'] * len(COLUMNS))}, false) "
    "on conflict (id) do update set "
    + ", ".join(f"{c}=excluded.{c}" for c in COLUMNS if c != "id")
    + ", editor_set=false, ingested_at=now()"
)


def cell_text(value):
    return "" if value is None else str(value).strip()


def header_field(value):
    text = cell_text(value).lstrip("\ufeff").casefold()
    text = re.sub(r"[_\-()\[\]:]+", " ", text)
    text = " ".join(text.split())
    text = re.sub(r"\s+(?:cm|см|սմ|amd|драм|դրամ|֏)$", "", text)
    if text in HEADER_FIELDS:
        return HEADER_FIELDS[text]
    # A sheet can repeat a photo heading or number its photo columns.
    base = re.sub(r"\s*\d+$", "", text).strip()
    return "photo" if HEADER_FIELDS.get(base) == "photo" else None


def parse_number(value):
    """A finite decimal, accepting sheet units and space-separated thousands."""
    if value is None or isinstance(value, bool):
        raise ValueError("missing or invalid number")
    text = cell_text(value).casefold()
    text = re.sub(r"^(?:amd|драм|դրամ|֏)\s*", "", text)
    text = re.sub(r"\s*(?:cm|см|սմ|amd|драм|դրամ|֏)$", "", text)
    text = re.sub(r"\s+", "", text).replace(",", ".")
    if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)", text):
        raise ValueError("unparseable number")
    try:
        number = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError("unparseable number") from exc
    if not number.is_finite():
        raise ValueError("unparseable number")
    return number


def split_urls(value):
    """Split the sheet's URL list, deduplicating in source order."""
    return list(dict.fromkeys(re.findall(r"https?://[^\s,]+", cell_text(value), re.I)))


class RowRejected(ValueError):
    def __init__(self, reasons):
        self.reasons = reasons
        super().__init__("; ".join(reasons))


def parse_row(pairs, slug, permission):
    """Parse (header, cell) pairs, preserving repeated photo columns."""
    pairs = list(pairs)
    fields = {}
    images = []
    for header, value in pairs:
        field = header_field(header)
        if field == "photo":
            images.extend(split_urls(value))
        elif field and cell_text(value):
            fields[field] = value
    reasons = []
    sku = cell_text(fields.get("sku"))
    if not sku:
        reasons.append("sku: missing")
    numbers = {}
    for field in ("width", "depth", "height", "price"):
        value = fields.get(field)
        if not cell_text(value):
            reasons.append(f"{field}: missing")
            continue
        try:
            number = parse_number(value)
        except ValueError:
            reasons.append(f"{field}: unparseable")
            continue
        if field == "price":
            if not 0 <= number <= 2147483647:
                reasons.append("price: outside 0–2147483647 AMD")
                continue
            numbers[field] = int(number.quantize(Decimal("1"), rounding=ROUND_HALF_UP))
        else:
            if not 5 <= number <= 400:
                reasons.append(f"{field}: outside 0.05–4 m")
                continue
            numbers[field] = float(number / 100)
    if reasons:
        raise RowRejected(reasons)
    size = [numbers[axis] for axis in ("width", "depth", "height")]
    name = cell_text(fields.get("name")) or None
    category = cell_text(fields.get("category")) or None
    images = list(dict.fromkeys(images))
    return dict(
        id=f"{slug}:{sku}", source=slug, source_id=sku, name=name, brand=None,
        description=cell_text(fields.get("url")) or None, product_type=category,
        kind=kind_of(product_type=(category or "").upper(), name=name),
        size_m=size, fit_size_m=size.copy(), size_status="estimated",
        size_evidence={"from": "shop sheet"}, listing_size_m=size.copy(),
        price=numbers["price"], currency="AMD", price_source="shop",
        color_text=[cell_text(fields["color"])] if "color" in fields else [],
        color_std=[], materials=[cell_text(fields["material"])] if "material" in fields else [],
        styles=[], keywords=[], main_image_url=images[0] if images else None,
        image_urls=images, glb_url=None, license=permission,
        raw=[[cell_text(h), cell_text(v)] for h, v in pairs], editor_set=False,
    )


def read_sheet(path):
    """Yield sheet row numbers and pairs without collapsing duplicate headings."""
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            sample = stream.read(8192)
            stream.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            reader = csv.reader(stream, dialect)
            headers = next(reader, [])
            for number, values in enumerate(reader, 2):
                if any(cell_text(v) for v in values):
                    yield number, list(zip(headers, values))
    elif path.suffix.lower() == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("XLSX requires openpyxl; run uv run --with openpyxl ingest_shop.py ...") from exc
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            reader = book.worksheets[0].iter_rows(values_only=True)
            headers = next(reader, [])
            for number, values in enumerate(reader, 2):
                if any(cell_text(v) for v in values):
                    yield number, list(zip(headers, values))
        finally:
            book.close()
    else:
        raise ValueError("--file must be CSV or XLSX")


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--shop", required=True, help="lowercase shop slug")
    parser.add_argument("--permission", required=True, help="written permission text, stored verbatim")
    parser.add_argument("--file", type=Path, required=True)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args(argv)
    if not re.fullmatch(r"[a-z0-9]+(?:[-_][a-z0-9]+)*", args.shop):
        parser.error("--shop must be a lowercase slug using letters, digits, hyphens or underscores")
    if not args.permission.strip():
        parser.error("--permission must contain written permission text")
    rows, rejected, reasons, kinds = [], 0, Counter(), Counter()
    try:
        for number, pairs in read_sheet(args.file):
            try:
                row = parse_row(pairs, args.shop, args.permission)
            except RowRejected as exc:
                rejected += 1
                reasons.update(exc.reasons)
                print(f"row {number}: rejected: {exc}")
                continue
            rows.append(row)
            kinds[row["kind"]] += 1
    except (OSError, ValueError) as exc:
        parser.error(str(exc))
    if not args.dry_run and rows:
        db_url = os.environ.get("VARPET_DB_URL")
        if not db_url:
            parser.error("VARPET_DB_URL is required unless --dry-run is used")
        with psycopg.connect(db_url) as conn:
            conn.execute((Path(__file__).parent / "schema.sql").read_text())
            with conn.cursor() as cur:
                cur.executemany(UPSERT, [
                    tuple(Jsonb(row[c]) if c in ("size_evidence", "raw") else row[c] for c in COLUMNS)
                    for row in rows
                ])
    print(f"{'dry-run (no DB): ' if args.dry_run else ''}loaded: {len(rows)}, rejected: {rejected}")
    print("rejected reasons:")
    for reason, count in sorted(reasons.items()):
        print(f"  {reason}: {count}")
    print("by kind:")
    for kind, count in sorted(kinds.items()):
        print(f"  {kind}: {count}")


if __name__ == "__main__":
    main()
